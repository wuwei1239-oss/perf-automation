# -*- coding: utf-8 -*-
"""
从本地 Excel 加载品种规则、车间人员（与 WPS 多维表表头对齐的「行字典」结构）。
候选文件名与 process_performance.py / 用户说明一致。
"""
from __future__ import annotations

import glob
import os
from typing import Dict, List

# 与 perf_engine.get_process_scores 使用的列名一致：首列为品种专属号
RULE_CODE_HEADERS = ("专属号（批号前缀）", "专属号", "批号前缀", "代码", "品种代码")

SHEET_CANDIDATES_RULES = ("Sheet1", "绩效规则", "品种规则", "规则库")


def _script_dir() -> str:
    return os.path.dirname(os.path.abspath(__file__))


def discover_rule_workbooks() -> List[str]:
    base = _script_dir()
    names = [
        "自动化绩效考核26.04.26.xlsx",
        "自动化绩效考核26.04.26.xls",
        "建立.xlsx",
        "绩效规则库.xlsx",
        "04月份绩效.xlsx",
    ]
    found: List[str] = []
    for n in names:
        p = os.path.join(base, n)
        if os.path.isfile(p):
            found.append(p)
    # 通配：自动化绩效考核*.xlsx
    for p in sorted(glob.glob(os.path.join(base, "自动化绩效考核*.xlsx"))):
        if p not in found:
            found.append(p)
    return found


def _pick_code_column(headers: List[str]) -> int:
    for i, h in enumerate(headers):
        hs = str(h).strip()
        for key in RULE_CODE_HEADERS:
            if key in hs or hs == key:
                return i
    return 0


def load_rules_from_workbook(path: str) -> Dict[str, dict]:
    try:
        import openpyxl
    except ImportError:
        return {}

    wb = openpyxl.load_workbook(path, data_only=True)
    best_ws = None
    for sn in SHEET_CANDIDATES_RULES:
        if sn in wb.sheetnames:
            best_ws = wb[sn]
            break
    if best_ws is None:
        best_ws = wb[wb.sheetnames[0]]

    headers: List[str] = []
    for c in range(1, best_ws.max_column + 1):
        v = best_ws.cell(row=1, column=c).value
        headers.append(str(v).strip() if v is not None else "")

    code_col = _pick_code_column(headers)
    rules: Dict[str, dict] = {}
    for r in range(2, best_ws.max_row + 1):
        raw_code = best_ws.cell(row=r, column=code_col + 1).value
        if raw_code is None or str(raw_code).strip() == "":
            continue
        code = str(raw_code).strip()
        row_data: dict = {}
        for c in range(1, best_ws.max_column + 1):
            v = best_ws.cell(row=r, column=c).value
            key = headers[c - 1]
            if not key:
                continue
            if v is None or v == "":
                continue
            row_data[key] = v
        if row_data:
            rules[code] = row_data
    return rules


def load_merged_rules_from_excel() -> Dict[str, dict]:
    """后出现的文件覆盖先出现的同专属号规则（优先更具体的文件名顺序）。"""
    merged: Dict[str, dict] = {}
    for path in discover_rule_workbooks():
        part = load_rules_from_workbook(path)
        merged.update(part)
    return merged


def discover_worker_workbooks() -> List[str]:
    base = _script_dir()
    out: List[str] = []
    for n in ("04月份绩效.xlsx", "车间人员.xlsx", "自动化绩效考核26.04.26.xlsx"):
        p = os.path.join(base, n)
        if os.path.isfile(p) and p not in out:
            out.append(p)
    return out


def _sheet_has_header(ws, header: str) -> bool:
    for r in range(1, min(5, ws.max_row + 1)):
        for c in range(1, min(ws.max_column + 1, 40)):
            v = ws.cell(row=r, column=c).value
            if v and str(v).strip() == header:
                return True
    return False


def load_workers_from_excel() -> List[str]:
    try:
        import openpyxl
    except ImportError:
        return []

    names: List[str] = []
    for path in discover_worker_workbooks():
        wb = openpyxl.load_workbook(path, data_only=True)
        for sn in wb.sheetnames:
            ws = wb[sn]
            if not _sheet_has_header(ws, "姓名"):
                continue
            # 定位「姓名」列
            name_col = None
            header_row = None
            for r in range(1, min(6, ws.max_row + 1)):
                for c in range(1, min(ws.max_column + 1, 60)):
                    v = ws.cell(row=r, column=c).value
                    if v and str(v).strip() == "姓名":
                        name_col = c
                        header_row = r
                        break
                if name_col:
                    break
            if not name_col or not header_row:
                continue
            for r in range(header_row + 1, ws.max_row + 1):
                v = ws.cell(row=r, column=name_col).value
                if v is None:
                    continue
                s = str(v).strip()
                if s and s not in names:
                    names.append(s)
        if names:
            break
    return sorted(names)


def summarize_templates_for_api() -> dict:
    """
    供 /api/status 使用：若存在 04月份绩效.xlsx / 绩效规则库.xlsx，返回表名与首行表头摘要
    （与 process_performance.parse_sheet 语义对齐：04 模板为「按日 sheet + 第1列姓名」）。
    """
    out: dict = {"04月份绩效": None, "绩效规则库": None}
    base = _script_dir()
    p04 = os.path.join(base, "04月份绩效.xlsx")
    pr = os.path.join(base, "绩效规则库.xlsx")
    try:
        import openpyxl
    except ImportError:
        return out
    for label, path in (("04月份绩效", p04), ("绩效规则库", pr)):
        if not os.path.isfile(path):
            continue
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            ws0 = wb[wb.sheetnames[0]]
            row1 = [
                str(ws0.cell(row=1, column=c).value or "").strip()
                for c in range(1, min(ws0.max_column + 1, 25))
            ]
            out[label] = {"path": os.path.basename(path), "sheet_names": wb.sheetnames[:15], "first_row_headers": row1}
        except Exception as exc:
            out[label] = {"error": str(exc)}
    return out
