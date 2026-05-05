# -*- coding: utf-8 -*-
import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

wb = openpyxl.load_workbook('绩效规则库.xlsx', data_only=True)
ws = wb['绩效规则']
print('=== 绩效规则表头 ===')
for c in range(1, ws.max_column+1):
    v = ws.cell(row=1, column=c).value
    print(f'  列{c}: {v}')

print()
print('=== 第一行数据（第2行）===')
for c in range(1, ws.max_column+1):
    v = ws.cell(row=2, column=c).value
    print(f'  列{c}: {v}')
