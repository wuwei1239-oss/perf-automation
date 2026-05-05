# -*- coding: utf-8 -*-
"""
建立.xlsx 规则标准化改造：总工分包干制
逻辑：
1. 读取建立.xlsx的Sheet1
2. 找到所有"基础分"列，提取数字
3. 找到对应"定员"列
4. 计算 总工分 = 基础分数字 × 定员
5. 替换基础分列的值为计算结果
6. 输出自动化绩效规则库_标准版.xlsx
"""

import openpyxl
import re
import copy

# 定义映射关系：基础分列 -> 对应定员列（列号，从1开始）
# 基于对表头的分析
BASE_SCORE_COLS = {
    5: 8,    # 配料基础分 -> 配料定员
    9: 11,   # 混合基础分 -> 混合定员
    12: 13,  # 预混基础分 -> 预混定员
    15: 16,  # 制粒基础分 -> 制粒定员
    19: 20,  # 压片基础分 -> 压片定员
    21: 22,  # 包衣基础分 -> 包衣定员
    23: 24,  # 内包基础分 -> 内包定员
    25: 26,  # 外包基础分 -> 外包定员
}

def extract_number(text):
    """从文本中提取数字，如 '70分/批/人' -> 70.0, '1.1分/批' -> 1.1"""
    if text is None:
        return None
    text = str(text).strip()
    # 使用正则提取数字（含小数）
    match = re.search(r'(\d+(?:\.\d+)?)', text)
    if match:
        return float(match.group(1))
    return None

def is_simple_score(text):
    """判断是否是简单的基础分文本（仅包含数字和单位，不含公式描述）"""
    if text is None:
        return False
    text = str(text).strip()
    # 简单模式：数字+分/批/人 或 数字+分/批
    if re.match(r'^\d+(?:\.\d+)?\u5206/\u6279(?:/\u4eba)?$', text):
        return True
    # 也接受只有数字的情况（已经是数字）
    if re.match(r'^\d+(?:\.\d+)?$', text):
        return True
    return False

def is_complex_formula(text):
    """判断是否包含公式/多条件（如 高速3分\\n双极6分）"""
    if text is None:
        return False
    text = str(text).strip()
    # 包含换行符或条件描述的视为复杂公式
    if '\n' in text:
        return True
    # 包含中文条件描述
    if re.search(r'[\u9ad8\u4f4e\u53cc\u5355\u901f]', text):
        return True
    return False

def main():
    print("=" * 60)
    print("建立.xlsx 规则标准化改造 - 总工分包干制")
    print("=" * 60)

    # 读取文件
    wb = openpyxl.load_workbook('建立.xlsx', data_only=True)
    ws = wb['Sheet1']

    print(f"\n读取完成: {ws.max_row} 行, {ws.max_column} 列")

    # 打印表头（用于确认）
    print("\n【表头确认】")
    header_map = {
        5: "配料基础分", 6: "配料大清分", 7: "过筛分", 8: "配料定员",
        9: "混合基础分", 10: "混合大清分", 11: "混合定员",
        12: "预混基础分", 13: "预混定员", 14: "预混大清分",
        15: "制粒基础分", 16: "制粒定员", 17: "干法制粒装机", 18: "制粒大清",
        19: "压片基础分", 20: "压片定员",
        21: "包衣基础分", 22: "包衣定员",
        23: "内包基础分", 24: "内包定员",
        25: "外包基础分", 26: "外包定员",
    }
    for r in range(1, 4):
        print(f"  行{r}:", end="")
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            if cell.value is not None:
                label = header_map.get(c, f"Col{c}")
                val = str(cell.value)[:30]
                print(f" [{label}={val}]", end="")
        print()

    # 处理数据行（从第2行开始，第1行是表头）
    stats = {"processed": 0, "skipped_no_number": 0, "skipped_complex": 0, "skipped_no_staffing": 0}

    for row_idx in range(2, ws.max_row + 1):
        row_processed = False
        for base_col, staff_col in BASE_SCORE_COLS.items():
            base_cell = ws.cell(row=row_idx, column=base_col)
            staff_cell = ws.cell(row=row_idx, column=staff_col)

            base_val = base_cell.value
            staff_val = staff_cell.value

            # 跳过空值
            if base_val is None:
                continue

            base_str = str(base_val).strip()

            # 跳过"无"或"-"
            if base_str in ['无', '-', '']:
                continue

            # 如果已经是数字（纯数字），也跳过（可能已经处理过了）
            if isinstance(base_val, (int, float)) and not isinstance(base_val, bool):
                continue

            # 跳过复杂公式（如内/外包带条件描述）
            if is_complex_formula(base_str):
                stats["skipped_complex"] += 1
                continue

            # 提取数字
            base_num = extract_number(base_str)
            if base_num is None:
                stats["skipped_no_number"] += 1
                continue

            # 处理定员
            staff_num = None
            if isinstance(staff_val, (int, float)) and not isinstance(staff_val, bool):
                staff_num = float(staff_val)
            else:
                staff_str = str(staff_val).strip() if staff_val else ''
                # 如果定员是"无"或非数字，跳过
                if staff_str in ['无', '-', ''] or is_complex_formula(staff_str):
                    stats["skipped_no_staffing"] += 1
                    continue
                staff_num = extract_number(staff_str)
                if staff_num is None:
                    stats["skipped_no_staffing"] += 1
                    continue

            # 计算总工分
            total_score = base_num * staff_num
            # 如果是整数则显示整数
            if total_score == int(total_score):
                total_score = int(total_score)
            else:
                total_score = round(total_score, 1)

            label = header_map.get(base_col, f"Col{base_col}")
            print(f"  行{row_idx} [{label}]: {base_str} × {staff_num} = {total_score}")
            base_cell.value = total_score
            row_processed = True

        if row_processed:
            stats["processed"] += 1

    # 输出统计
    print(f"\n【处理统计】")
    print(f"  处理行数: {stats['processed']}")
    print(f"  跳过（无法提取数字）: {stats['skipped_no_number']}")
    print(f"  跳过（复杂公式）: {stats['skipped_complex']}")
    print(f"  跳过（无定员/定员复杂）: {stats['skipped_no_staffing']}")

    # 保存文件
    output_path = '自动化绩效规则库_标准版.xlsx'
    wb.save(output_path)
    print(f"\n[OK] 已生成: {output_path}")

    # 验证：读取前几行确认结果
    print("\n【验证 - 处理后的前5行数据】")
    wb2 = openpyxl.load_workbook(output_path)
    ws2 = wb2['Sheet1']
    for r in range(1, min(6, ws2.max_row + 1)):
        print(f"  行{r}:", end="")
        for c in [5, 8, 9, 11, 15, 16, 19, 20, 21, 22, 23, 24, 25, 26]:
            cell = ws2.cell(row=r, column=c)
            if cell.value is not None:
                label = header_map.get(c, f"Col{c}")
                print(f" [{label}={cell.value}]", end="")
        print()

if __name__ == '__main__':
    main()
