# -*- coding: utf-8 -*-
"""Helper to write v5 file"""
import sys, json, io, re, time, requests, os, openpyxl
from datetime import datetime
from collections import defaultdict

HOUR_SCORE = 20

# 列索引
COL = {
    '专属号':1,'配料基础分':5,'配料大清分':6,'过筛分':7,'配料定员':8,
    '混合基础分':9,'混合大清分':10,'混合定员':11,
    '预混基础分':12,'预混定员':13,'预混大清分':14,
    '制粒基础分':15,'制粒定员':16,'制粒装机':17,'制粒大清':18,
    '压片基础分':19,'压片定员':20,'压片77冲':21,'压片65冲':22,'压片40冲':23,
    '包衣基础分':24,'包衣定员':25,'包衣大清分':26,
    '内包基础分':27,'内包定员':28,'内包大清分':29,
    '外包基础分':30,'外包定员':31,'外包大清分':32,'捡药分':33,
}

PROCESS_COLS = {
    '配料':(5,6,8),'过筛':(7,None,None),'混合':(9,10,11),
    '预混':(12,14,13),'制粒':(15,18,16,17),
    '压片':(19,None,20,21,22,23),'包衣':(24,26,25),
    '内包':(27,29,28),'外包':(30,32,31),
}

OPTIONS = {
    '配料':['过筛','缺员','质量','转片子'],'搅拌':['缺员','质量','转片子'],
    '过筛':['缺员','质量','转片子'],'混合':['缺员','质量','转片子'],
    '预混':['缺员','质量','转片子'],
    '制粒':['干法制粒装机','缺员','质量','转片子'],
    '压片':['压片冲型','缺员','质量','转片子'],
    '包衣':['缺员','质量','转片子'],
    '内包':['大清分','缺员','质量','转片子','模具更换','模具确认','捡药','粉碎'],
    '外包':['大清分','缺员','质量','转片子','模具更换','模具确认','捡药','粉碎'],
}

def n(v,d=0):
    if v is None: return d
    if isinstance(v,(int,float)): return float(v)
    try: return float(str(v).strip())
    except: return d

def sv(v):
    if v is None: return ''
    if isinstance(v,(int,float)): return str(int(v)) if v==int(v) else str(v)
    return str(v)

def pc(v):
    if v is None: return 0
    if isinstance(v,(int,float)): return float(v)
    s=str(v).strip()
    if not s or s in ['--','-','无','']: return 0
    m=re.search(r'(\d+(?:\.\d+)?)',s)
    return float(m.group(1)) if m else 0

def pd(v):
    if v is None: return None
    if isinstance(v,(int,float)): return int(v) if int(v)>0 else None
    s=str(v).strip()
    if not s or s in ['--','-','无',''] or '无定员' in s or '几人生产' in s: return None
    m=re.search(r'(\d+)',s)
    return int(m.group(1)) if m else None

def parse_batch(batch_no,known=None):
    if not batch_no: return None,1
    s=str(batch_no).strip()
    pcnt=0
    m=re.search(r'[（(](\d+)\s*批[)）]',s)
    if m:
        pcnt=int(m.group(1))
        s=re.sub(r'[（(].*?[)）]','',s)
    s=s.replace('批','')
    parts=re.split(r'[,，+、\s]+',s)
    cnt=pcnt if pcnt>0 else 0; prefix=None
    for part in parts:
        part=part.strip()
        if not part: continue
        if '-' in part:
            left,right=part.split('-')[0].strip(),part.split('-')[1].strip()
            ld=re.sub(r'\D','',left); rd=re.sub(r'\D','',right)
            if ld and rd:
                try:
                    rv=int(rd); lv=int(ld[-len(rd):]) if len(ld)>=len(rd) else int(ld)
                    cnt+=(rv-lv+1) if rv>=lv else 1
                except: cnt+=1
            else: cnt+=1
            if prefix is None: prefix=ex(ld,known)
        else:
            cnt+=1
            if prefix is None: prefix=ex(re.sub(r'\D','',part),known)
    return prefix,cnt if cnt>0 else 1

def ex(digits,known=None):
    if not digits: return None
    digits=str(digits)
    for l in [4,3,2]:
        if len(digits)>=l:
            p=digits[:l]
            if known and p in known: return p
    return digits[:3] if len(digits)>=3 else (digits[:2] if len(digits)>=2 else None)

def get_scores(row,proc,box=0,chong=None):
    cols=PROCESS_COLS.get(proc,None)
    r={'base':0,'daqing':0,'ding':None,'zhuangji':0,'guoshao':0}
    if not cols: return r
    bi,dqi,di=cols[0],cols[1],cols[2]
    if bi<=len(row): r['base']=pc(row[bi-1])
    if dqi and dqi<=len(row): r['daqing']=pc(row[dqi-1])
    if di and di<=len(row): r['ding']=pd(row[di-1])
    if proc=='配料':
        r['guoshao']=pc(row[6]) if len(row)>6 else 0
    if proc=='制粒' and len(cols)>=4:
        zji=cols[3]
        if zji<=len(row): r['zhuangji']=pc(row[zji-1])
    if proc=='压片' and chong:
        cm={'77':21,'65':22,'40':23}
        if chong in cm:
            ci=cm[chong]-1
            if ci<len(row): r['zhuangji']=pc(row[ci])
    return r

def calc(records,rules,prefixes):
    results=[]
    for rec in records:
        proc=sv(rec.get('工序',''))
        batch_no=sv(rec.get('批号',''))
        box=n(rec.get('批次量(箱)',0)); hours=n(rec.get('工时',8))
        has_dq=rec.get('是否有大清分','否')=='是'
        has_gs=rec.get('是否有过筛分','否')=='是'
        has_zj=rec.get('是否有干法制粒装机','否')=='是'
        chong=sv(rec.get('压片冲型',''))
        qual=n(rec.get('质量总分',0)); zpn=n(rec.get('转片子批数',0))
        jyn=n(rec.get('捡药筐数',0)); qf=n(rec.get('缺员分',0))
        hqf=n(rec.get('小时缺员分',0)); fensui=n(rec.get('粉碎分',0))
        
        prefix,batch_n=parse_batch(batch_no,prefixes)
        rule_row=rules.get(prefix,[])
        is_exp=not rule_row
        scores=get_scores(rule_row,proc,box,chong)
        
        names=[s.strip() for s in sv(rec.get('人员姓名','')).split(',') if s.strip()]
        css=[s.strip() for s in sv(rec.get('人员系数','')).split(',') if s.strip()]
        coeffs=[]
        for i in range(len(names)):
            try: coeffs.append(float(css[i]) if i<len(css) else 1.0)
            except: coeffs.append(1.0)
        if not names: continue
        
        nn=len(names)
        ding=scores['ding'] if scores['ding'] else nn
        
        if is_exp: bpp=(hours*HOUR_SCORE)/nn
        else: bpp=(scores['base']*batch_n)/ding if ding>0 else (scores['base']*batch_n)/nn
        
        qns=[s.strip() for s in sv(rec.get('缺员人员','')).split(',') if s.strip()]
        qc=0
        if qf>0: qc=qf/nn
        elif qns and ding>nn and not is_exp:
            qc=(scores['base']*batch_n/ding)*len(qns)/nn
        
        pp_dq=(scores['daqing']/nn) if (has_dq and scores['daqing']>0) else 0
        pp_gs=(scores['guoshao']/nn) if (has_gs and scores['guoshao']>0) else 0
        pp_zj=(scores['zhuangji']/nn) if (has_zj and scores['zhuangji']>0) else 0
        pp_q=qual/nn if qual>0 else 0
        pp_hq=hqf/nn if hqf>0 else 0
        pp_fs=fensui/nn if fensui>0 else 0
        
        zns=[s.strip() for s in sv(rec.get('转片子人员','')).split(',') if s.strip()]
        pp_zp=(5*zpn)/len(zns) if (zpn>0 and zns) else 0
        
        jns=[s.strip() for s in sv(rec.get('捡药人员','')).split(',') if s.strip()]
        jt=0
        if jyn>0 and jns and rule_row:
            jpv=pc(rule_row[32]) if len(rule_row)>32 else 0
            jt=(jyn*jpv)/len(jns) if jpv>0 else 0
        
        mgc=[s.strip() for s in sv(rec.get('更换模具人员','')).split(',') if s.strip()]
        mgcf=[s.strip() for s in sv(rec.get('模具确认人员','')).split(',') if s.strip()]
        
        qhs=sv(rec.get('缺员小时数',''))
        qhl=[n(h) for h in qhs.split(',') if h.strip()]
        ded={}
        for qn,qh in zip(qns,qhl):
            if qh>0: ded[qn]=qh*HOUR_SCORE
        
        for i,name in enumerate(names):
            c=coeffs[i]
            wc=bpp*c
            nc=pp_dq+pp_gs+pp_zj+pp_q+pp_hq+pp_fs+qc
            if name in zns: nc+=pp_zp
            if name in jns: nc+=jt
            if name in mgc: nc+=30
            if name in mgcf: nc+=10
            d=ded.get(name,0)
            total=wc+nc-d
            
            parts=[]
            if not is_exp and bpp>0: parts.append(f"{scores['base']}x{batch_n}批/{ding}人x{c}")
            if is_exp: parts.append(f"实验{hours}hx20分x{c}")
            if pp_dq>0: parts.append(f"大清+{pp_dq:.1f}")
            if pp_gs>0: parts.append(f"过筛+{pp_gs:.1f}")
            if pp_zj>0: parts.append(f"装机+{pp_zj:.1f}")
            if qc>0: parts.append(f"缺员补+{qc:.1f}")
            if pp_q>0: parts.append(f"质量+{pp_q:.1f}")
            if pp_hq>0: parts.append(f"时缺+{pp_hq:.1f}")
            if pp_fs>0: parts.append(f"粉碎+{pp_fs:.1f}")
            if name in zns and pp_zp>0: parts.append(f"转片+{pp_zp:.1f}")
            if name in mgc: parts.append("模具换+30")
            if name in mgcf: parts.append("模具确+10")
            if name in jns and jt>0: parts.append(f"捡药+{jt:.1f}")
            if d>0: parts.append(f"扣{d:.0f}")
            
            results.append({
                '日期':rec.get('生产日期',''),'姓名':name,'工序':proc,'系数':c,
                '批号':batch_no,'箱数':box,'批数':batch_n,'工时':hours,
                '基础分':round(bpp,2),'大清':round(pp_dq,2),'过筛':round(pp_gs,2),
                '装机':round(pp_zj,2),'缺员补':round(qc,2),'时缺':round(pp_hq,2),
                '质量':round(pp_q,2),'粉碎':round(pp_fs,2),
                '转片':round(pp_zp if name in zns else 0,2),
                '模具换':30 if name in mgc else 0,
                '模具确':10 if name in mgcf else 0,
                '捡药':round(jt if name in jns else 0,2),'扣':round(d,2),
                '系数分':round(wc,2),'加减分':round(nc,2),
                '总分':round(total,2),'备注':'; '.join(parts),
            })
    return results

def monthly(res):
    s=defaultdict(lambda:{'天':set(),'分':0,'工序':set()})
    for r in res:
        s[r['姓名']]['天'].add(r['日期'])
        s[r['姓名']]['分']+=r['总分']
        s[r['姓名']]['工序'].add(r['工序'])
    out=[]
    for name,info in sorted(s.items(),key=lambda x:x[1]['分'],reverse=True):
        d=len(info['天'])
        out.append({'姓名':name,'出勤天数':d,'总得分':round(info['分'],2),
                    '日均分':round(info['分']/d,2) if d>0 else 0,
                    '工序':'、'.join(sorted(info['工序']))})
    return out

def read_excel(path='绩效规则库.xlsx'):
    wb=openpyxl.load_workbook(path,data_only=True)
    ws=wb['绩效规则']
    rules={}
    for r in range(2,ws.max_row+1):
        code=ws.cell(row=r,column=1).value
        if code is None: continue
        code=str(code).strip()
        if not code: continue
        row=[]
        for c in range(1,34): row.append(ws.cell(row=r,column=c).value)
        rules[code]=row
    return rules,set(rules.keys())

def run_local():
    rules,prefixes=read_excel()
    print(f"规则库: {len(rules)}条 keys: {list(rules.keys())}")
    
    tests=[
        {'日期':'2026-04-28','工序':'配料','批号':'02260401-03','箱数':0,'工时':8,
         '人员':'郭红杰,王磊,刘秀霞','系数':'1.22,1.0,1.0',
         '大清':'是','过筛':'否','装机':'否','冲型':'','缺员分':0,'时缺分':0,
         '质量分':0,'转片批':0,'转片人':'','模具换':'','模具确':'',
         '捡药筐':0,'捡药人':'','缺员人':'','缺员时':'','粉碎':0},
        {'日期':'2026-04-28','工序':'压片','批号':'02260404','箱数':0,'工时':8,
         '人员':'刘好明,唐雪飞,杨芬,任宏娟','系数':'1.22,1.10,1.0,1.0',
         '大清':'否','过筛':'否','装机':'否','冲型':'65','缺员分':0,'时缺分':0,
         '质量分':20,'转片批':0,'转片人':'','模具换':'','模具确':'',
         '捡药筐':0,'捡药人':'','缺员人':'','缺员时':'','粉碎':0},
        {'日期':'2026-04-28','工序':'外包','批号':'36260306批,36260307批','箱数':409,'工时':8,
         '人员':'刘金浩,石法智,杜娟,王伟伟,宋泉泉,王桂霞,吴艳,王东松,姚娟',
         '系数':'1.20,1.10,1.06,1.06,1.06,1.06,1.06,1.0,1.0',
         '大清':'是','过筛':'否','装机':'否','冲型':'','缺员分':0,'时缺分':0,
         '质量分':0,'转片批':0,'转片人':'','模具换':'','模具确':'',
         '捡药筐':30,'捡药人':'刘金浩,宋泉泉,王桂霞,吴艳','缺员人':'','缺员时':'','粉碎':2},
        {'日期':'2026-04-28','工序':'内包','批号':'02260401','箱数':240,'工时':8,
         '人员':'李勇,张勇攀,苗广秀,杜曼','系数':'1.20,1.08,1.08,1.0',
         '大清':'是','过筛':'否','装机':'否','冲型':'','缺员分':0,'时缺分':0,
         '质量分':20,'转片批':0,'转片人':'','模具换':'李勇','模具确':'张勇攀',
         '捡药筐':0,'捡药人':'','缺员人':'','缺员时':'','粉碎':0},
        {'日期':'2026-04-28','工序':'压片','批号':'260401','箱数':0,'工时':2,
         '人员':'刘好明,唐雪飞','系数':'1.22,1.10',
         '大清':'否','过筛':'否','装机':'否','冲型':'','缺员分':0,'时缺分':0,
         '质量分':10,'转片批':0,'转片人':'','模具换':'','模具确':'',
         '捡药筐':0,'捡药人':'','缺员人':'','缺员时':'','粉碎':0},
    ]
    
    fmt=[]
    for t in tests:
        fmt.append({'生产日期':t['日期'],'工序':t['工序'],'批号':t['批号'],
            '批次量(箱)':t['箱数'],'工时':t['工时'],
            '人员姓名':t['人员'],'人员系数':t['系数'],
            '是否有大清分':t['大清'],'是否有过筛分':t['过筛'],
            '是否有干法制粒装机':t['装机'],'压片冲型':t['冲型'],
            '缺员分':t['缺员分'],'小时缺员分':t['时缺分'],
            '质量总分':t['质量分'],'转片子批数':t['转片批'],
            '转片子人员':t['转片人'],'更换模具人员':t['模具换'],
            '模具确认人员':t['模具确'],'捡药筐数':t['捡药筐'],
            '捡药人员':t['捡药人'],'缺员人员':t['缺员人'],
            '缺员小时数':t['缺员时'],'粉碎分':t['粉碎']})
    
    res=calc(fmt,rules,prefixes)
    
    wb=openpyxl.Workbook()
    ws=wb.active; ws.title='绩效明细'
    hs=['日期','姓名','工序','系数','批号','箱数','批数','工时',
        '基础分','大清','过筛','装机','缺员补','时缺','质量','粉碎',
        '转片','模具换','模具确','捡药','扣','系数分','加减分','总分','备注']
    
    from openpyxl.styles import Font,Alignment,Border,Side,PatternFill
    hf=Font(name='微软雅黑',size=10,bold=True,color='FFFFFF')
    hfill=PatternFill(start_color='4472C4',end_color='4472C4',fill_type='solid')
    ha=Alignment(horizontal='center',vertical='center',wrap_text=True)
    bd=Border(left=Side(style='thin'),right=Side(style='thin'),
              top=Side(style='thin'),bottom=Side(style='thin'))
    
    for c,h in enumerate(hs,1):
        cell=ws.cell(row=1,column=c,value=h)
        cell.font,cell.fill,cell.alignment,cell.border=hf,hfill,ha,bd
    
    df=Font(name='微软雅黑',size=9)
    da=Alignment(horizontal='center',vertical='center',wrap_text=True)
    
    for i,r in enumerate(res,2):
        for c,h in enumerate(hs,1):
            v=r.get(h,'')
            cell=ws.cell(row=i,column=c,value=v)
            cell.font,cell.alignment,cell.border=df,da,bd
    
    widths=[12,8,7,6,14,6,5,5,7,5,5,5,6,6,5,5,5,5,5,5,5,7,7,8,40]
    for i,w in enumerate(widths,1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    
    ws2=wb.create_sheet('月度汇总')
    for c,h in enumerate(['姓名','出勤天数','总得分','日均分','工序'],1):
        cell=ws2.cell(row=1,column=c,value=h)
        cell.font,cell.fill,cell.alignment,cell.border=hf,hfill,ha,bd
    
    sm=monthly(res)
    for i,m in enumerate(sm,2):
        for c,v in enumerate([m['姓名'],m['出勤天数'],m['总得分'],m['日均分'],m['工序']],1):
            cell=ws2.cell(row=i,column=c,value=v)
            cell.font,cell.alignment,cell.border=df,da,bd
    
    wb.save('绩效测试v5.xlsx')
    print(f"\n结果: {len(res)}条, 月汇总{len(sm)}人")
    print(f"\n{'姓名':<8} {'天':<4} {'总得分':<8} {'日均':<6} {'工序'}")
    for m in sm:
        print(f"{m['姓名']:<8} {m['出勤天数']:<4} {m['总得分']:<8.2f} {m['日均分']:<6.2f} {m['工序']}")
    
    print("\n详细:")
    for r in res:
        print(f"{r['日期']} {r['姓名']:<6} {r['工序']:<4} 系数{r['系数']:<4} {r['总分']:>7.2f}分 | {r['备注']}")

if __name__=='__main__':
    run_local()
